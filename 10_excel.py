"""Costruisce la cartella di lavoro: clienti-che-tornano.xlsx

Legge sito/dati.json, gli stessi numeri della pagina e di RISULTATI.md. Non
rifa' nessun conto: se una cifra qui non torna con i documenti, il colpevole e'
uno solo ed e' 08_dati_pagina.py.

Il foglio del pareggio contiene FORMULE VERE, non valori incollati: chi apre il
file cambia il margine nella sua cella e la soglia si ricalcola. Era la
richiesta di DOMANDA.md — «il margine resta un parametro dichiarato, non un
risultato» — e in un foglio di calcolo si mantiene meglio che altrove.

Uso:  python 10_excel.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

QUI = Path(__file__).parent
DATI = QUI / "sito" / "dati.json"
USCITA = QUI / "clienti-che-tornano.xlsx"

MESI = ["gen", "feb", "mar", "apr", "mag", "giu",
        "lug", "ago", "set", "ott", "nov", "dic"]

INCHIOSTRO = "FF15181B"
TENUE = "FF5D656C"
OTTONE = "FF8A6414"
CARTA = "FFF2F3F4"
INPUT = "FFFDF3D6"

T_TITOLO = Font(name="Calibri", size=17, bold=True, color=INCHIOSTRO)
T_H2 = Font(name="Calibri", size=12, bold=True, color=OTTONE)
T_TESTA = Font(name="Calibri", size=9, bold=True, color=TENUE)
T_NORM = Font(name="Calibri", size=10.5, color=INCHIOSTRO)
T_TENUE = Font(name="Calibri", size=9.5, color=TENUE)
T_FORTE = Font(name="Calibri", size=10.5, bold=True, color=INCHIOSTRO)
T_GROSSO = Font(name="Calibri", size=20, bold=True, color=OTTONE)

SOTTILE = Side(style="thin", color="FFD6DADE")
BORDO = Border(bottom=SOTTILE)
BOX = Border(left=Side(style="medium", color=OTTONE), right=Side(style="medium", color=OTTONE),
             top=Side(style="medium", color=OTTONE), bottom=Side(style="medium", color=OTTONE))


def scrivi(ws, cella, valore, font=T_NORM, formato=None, riemp=None,
           allinea=None, bordo=None):
    c = ws[cella]
    c.value = valore
    c.font = font
    if formato:
        c.number_format = formato
    if riemp:
        c.fill = PatternFill("solid", fgColor=riemp)
    if allinea:
        c.alignment = allinea
    if bordo:
        c.border = bordo
    return c


def testo(ws, riga, righe, col="A", font=T_NORM):
    for t in righe:
        if t:
            scrivi(ws, f"{col}{riga}", t, font)
        riga += 1
    return riga


# ═════════════════════════════════════════════════════════════════════════
def foglio_leggimi(wb, d):
    ws = wb.active
    ws.title = "Leggimi"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 104

    scrivi(ws, "B2", "Clienti che tornano", T_TITOLO)
    scrivi(ws, "B3", "Online Retail II — negozio online britannico di articoli da regalo, "
                     "dicembre 2009 – dicembre 2011, 1.067.371 righe.", T_TENUE)

    r = 5
    scrivi(ws, f"B{r}", "LA DOMANDA", T_H2); r += 1
    r = testo(ws, r, [
        "Quanto vale un cliente che torna, e conviene di piu' trattenerlo o acquisirne "
        "uno nuovo?",
        "Scritta prima di aprire SQL. Il committente immaginabile: chi decide il budget "
        "marketing di un negozio online.",
        "",
    ])

    v = d["valore"]
    scrivi(ws, f"B{r}", "LA RISPOSTA", T_H2); r += 1
    r = testo(ws, r, [
        f"Fra i clienti entrati nel 2010, chi arriva al secondo ordine spende nel primo "
        f"anno {v['torna']['media']:,} £ contro {v['non_torna']['media']:,} £."
        .replace(",", "."),
        f"Confrontando solo clienti che al primo acquisto erano uguali, il divario resta "
        f"{v['appaiata']:,} £ (intervallo {v['appaiata_ic'][0]:,}–{v['appaiata_ic'][1]:,}): "
        f"ne sopravvive l'{round(v['appaiata']/v['grezza']*100)}%.".replace(",", "."),
        "Riattivare conviene finche' costa meno di  margine × 1.125 £.  "
        "Il foglio «Valore e pareggio» lo calcola col tuo margine.",
        "",
    ])

    scrivi(ws, f"B{r}", "E LA COSA CHE NON ERA IN PROGRAMMA", T_H2); r += 1
    r = testo(ws, r, [
        "La retention di questo negozio non e' una curva di abbandono: e' un calendario.",
        f"Il mese dell'anno spiega il {d['spiegata']['calendario']}% della variazione fra "
        f"le celle, l'eta' del cliente solo il {d['spiegata']['eta']}%. Novembre 25%, "
        f"gennaio 11%.",
        "Leggere la matrice da sinistra a destra vuol dire leggere il Natale e chiamarlo "
        "fedelta'. I due fogli «Matrice» mostrano le stesse celle nei due allineamenti: "
        "confrontali.",
        "",
    ])

    scrivi(ws, f"B{r}", "COSA QUESTI NUMERI NON DICONO", T_H2); r += 1
    r = testo(ws, r, [
        "• Non dicono se trattenere convenga piu' che acquisire: manca il costo di "
        "acquisizione. Dicono a che prezzo le due cose si equivalgono.",
        "• Misurano ricavo, non profitto: manca il costo del venduto. Il margine e' un "
        "parametro dichiarato, non un risultato.",
        f"• Il 22,8% delle righe non ha un identificativo cliente. Il tasso di ritorno "
        f"vero sta fra il {d['anonimi']['quota_bassa']}% e il {d['anonimi']['quota_alta']}%: "
        f"si scrive «tre clienti IDENTIFICATI su quattro tornano».",
        "• Non dimostrano una causa. «I clienti che tornano valgono N volte» e' vero; "
        "«farli tornare li fa valere N volte» richiederebbe un esperimento.",
        "• I dati sono del 2009-2011. Il metodo si trasferisce, i numeri no.",
        "",
    ], font=T_TENUE)

    scrivi(ws, f"B{r}", "COME E' STATO FATTO", T_H2); r += 1
    testo(ws, r, [
        "Python e SQL (MySQL 8, funzioni finestra). Le pulizie, il modello, le verifiche e "
        "i limiti stanno nel repo, un documento per ognuno.",
        "Le cifre di questo file escono dallo stesso file di dati della pagina web, e sono "
        "ricontrollate da uno script che rifa' 63 conti dalle tabelle di base.",
    ], font=T_TENUE)


# ═════════════════════════════════════════════════════════════════════════
def foglio_matrice(wb, d, per_calendario: bool):
    nome = "Matrice per calendario" if per_calendario else "Matrice per età"
    ws = wb.create_sheet(nome)
    ws.sheet_view.showGridLines = False
    c0 = d["coorti"][0]
    off = lambda c: (c["anno"] - c0["anno"]) * 12 + (c["mese"] - c0["mese"])
    n_col = max(off(c) + len(c["celle"]) for c in d["coorti"]) if per_calendario \
        else max(len(c["celle"]) for c in d["coorti"])

    scrivi(ws, "A1", "Matrice per mese del calendario" if per_calendario
           else "Matrice per età del cliente", T_TITOLO)
    scrivi(ws, "A2",
           "Ogni riga è un mese di clienti nuovi. Le celle dicono quanti di loro hanno "
           "ordinato in quel mese, su cento." if not per_calendario else
           "Le STESSE celle del foglio precedente, spostate: ogni riga parte dal suo mese "
           "vero. Ora si legge dall'alto in basso.", T_TENUE)
    scrivi(ws, "A3",
           "Si legge da sinistra a destra, e sembra una discesa. Il foglio dopo mostra "
           "perché non lo è." if not per_calendario else
           "Le colonne si accendono e si spengono tutte insieme: novembre acceso in ogni "
           "riga, gennaio spento in ogni riga. Non è l'età del cliente a muovere il "
           "numero, è il mese dell'anno.", T_TENUE)

    R0 = 5
    scrivi(ws, f"A{R0}", "coorte", T_TESTA)
    scrivi(ws, f"B{R0}", "clienti", T_TESTA, allinea=Alignment(horizontal="right"))
    for i in range(n_col):
        col = get_column_letter(3 + i)
        et = (f"{MESI[(c0['mese'] - 1 + i) % 12]} "
              f"{str(c0['anno'] + (c0['mese'] - 1 + i) // 12)[2:]}"
              if per_calendario else ("0" if i == 0 else f"+{i}"))
        scrivi(ws, f"{col}{R0}", et, T_TESTA, allinea=Alignment(horizontal="center"))
        ws.column_dimensions[col].width = 6.5
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 8

    for j, c in enumerate(d["coorti"]):
        r = R0 + 1 + j
        scrivi(ws, f"A{r}", c["coorte"], T_NORM)
        scrivi(ws, f"B{r}", c["entrati"], T_NORM, "#,##0")
        base = off(c) if per_calendario else 0
        for i, val in enumerate(c["celle"]):
            if val is None:
                continue
            col = get_column_letter(3 + base + i)
            cel = scrivi(ws, f"{col}{r}", val / 100, T_NORM, "0%",
                         allinea=Alignment(horizontal="center"))
            if i == 0:
                cel.font = Font(name="Calibri", size=9, color=TENUE)

    ultima = R0 + len(d["coorti"])
    zona = f"D{R0 + 1}:{get_column_letter(2 + n_col)}{ultima}"
    # La scala diverge intorno alla retention media: sotto la media le celle
    # leggono fredde, sopra calde. Non e' decorazione, e' la differenza che
    # questi due fogli devono far vedere.
    ws.conditional_formatting.add(zona, ColorScaleRule(
        start_type="num", start_value=0.05, start_color="FFA8C0D4",
        mid_type="num", mid_value=0.185, mid_color="FFF4EFE3",
        end_type="num", end_value=0.40, end_color="FFB8801F"))

    ws.freeze_panes = f"C{R0 + 1}"
    r = ultima + 2
    testo(ws, r, [
        "La colonna 0 è il mese di ingresso: vale 100 per definizione, non è un risultato.",
        "Le celle vuote a destra non sono zeri: sono mesi che il calendario non ha ancora "
        "raggiunto quando i dati finiscono (9 dicembre 2011).",
    ], font=T_TENUE)
    if per_calendario:
        scrivi(ws, f"A{r + 3}",
               "Le colonne di questo foglio sono mesi veri: ogni colonna contiene un solo "
               "mese del calendario, per tutte le coorti. Verificato contando.", T_TENUE)


# ═════════════════════════════════════════════════════════════════════════
def foglio_riacquisto(wb, d):
    ws = wb.create_sheet("Riacquisto")
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 14), ("B", 11), ("C", 11), ("D", 26), ("E", 13), ("F", 13)):
        ws.column_dimensions[col].width = w

    scrivi(ws, "A1", "Il riacquisto", T_TITOLO)
    scrivi(ws, "A2", f"Sui {d['campione']['clienti']:,} clienti entrati nel 2010 che hanno "
                     f"avuto 365 giorni interi davanti. Finestra uguale per tutti: se no si "
                     f"misura quanto tempo ha avuto il cliente, non il cliente."
           .replace(",", "."), T_TENUE)

    R0 = 4
    for i, t in enumerate(["gradino", "clienti", "quota", "sopravvive dal precedente"]):
        scrivi(ws, f"{get_column_letter(1+i)}{R0}", t, T_TESTA, bordo=BORDO,
               allinea=Alignment(horizontal="right" if i else "left"))
    for j, g in enumerate(d["funnel"]):
        r = R0 + 1 + j
        scrivi(ws, f"A{r}", f"{g['gradino']}º ordine", T_NORM)
        scrivi(ws, f"B{r}", g["clienti"], T_NORM, "#,##0")
        scrivi(ws, f"C{r}", g["quota"] / 100, T_NORM, "0,0%")
        if g["sopravvive"]:
            scrivi(ws, f"D{r}", g["sopravvive"] / 100, T_FORTE, "0,0%")

    r = R0 + len(d["funnel"]) + 2
    scrivi(ws, f"A{r}", f"La colonna di destra non si muove: {d['sopravvivenza_media']}% a "
                        f"ogni gradino.", T_FORTE); r += 1
    r = testo(ws, r, [
        "Arrivare al secondo ordine non rende un cliente più fedele di quanto fosse. "
        "«Portali al secondo acquisto e sono tuoi» qui è falso.",
        "",
        "Contando tutta la vita di ogni cliente veniva 72%, poi 78%, 80%, 81% — una fedeltà "
        "che cresce. Ma le vite non sono lunghe uguali: chi è entrato a gennaio 2010 ha "
        "avuto due anni per arrivare all'ottavo ordine, chi è entrato a settembre 2011 ne "
        "ha avuti tre mesi. Quella salita era il tempo in più di chi era arrivato in fondo.",
        "",
    ], font=T_TENUE)

    scrivi(ws, f"A{r}", "IL CONTROLLO", T_H2); r += 1
    scrivi(ws, f"A{r}", "Se la sopravvivenza è costante, il numero di ordini deve seguire "
                        "una geometrica. Il parametro è stimato dagli stessi dati, quindi "
                        "è una riscrittura coerente, non una prova indipendente.", T_TENUE)
    r += 2
    for i, t in enumerate(["ordini", "attesa", "osservata"]):
        scrivi(ws, f"{get_column_letter(1+i)}{r}", t, T_TESTA, bordo=BORDO,
               allinea=Alignment(horizontal="right" if i else "left"))
    for j, g in enumerate(d["geometrica"]):
        rr = r + 1 + j
        scrivi(ws, f"A{rr}", g["ordini"], T_NORM)
        scrivi(ws, f"B{rr}", g["attesa"] / 100, T_NORM, "0,0%")
        scrivi(ws, f"C{rr}", g["osservata"] / 100, T_NORM, "0,0%")

    # il tempo al secondo ordine
    r = r + len(d["geometrica"]) + 3
    scrivi(ws, f"A{r}", "QUANTO TEMPO PASSA FRA PRIMO E SECONDO ORDINE", T_H2); r += 1
    p = d["tempo_secondo"]["percentili"]
    scrivi(ws, f"A{r}", f"Mediana {p['50']} giorni. Un quarto torna entro {p['25']}, "
                        f"tre quarti entro {p['75']}, nove su dieci entro {p['90']}.",
           T_NORM); r += 1
    scrivi(ws, f"A{r}", "Una soglia di silenzio oltre la quale il cliente è perso NON "
                        "esiste: la curva cala in modo regolare e non precipita mai. "
                        "Dichiararne una vorrebbe dire inventarla.", T_TENUE)
    r += 2
    for i, t in enumerate(["giorni", "già tornati"]):
        scrivi(ws, f"{get_column_letter(1+i)}{r}", t, T_TESTA, bordo=BORDO,
               allinea=Alignment(horizontal="right" if i else "left"))
    for j, punto in enumerate(d["tempo_secondo"]["cumulata"][::3]):
        rr = r + 1 + j
        scrivi(ws, f"A{rr}", punto["giorni"], T_NORM)
        scrivi(ws, f"B{rr}", punto["quota"] / 100, T_NORM, "0,0%")
    scrivi(ws, f"D{r}", "Percentuali di chi torna ENTRO L'ANNO: la finestra taglia a 365 "
                        "giorni, quindi non sono percentuali di «chi tornerà» mai.", T_TENUE)


# ═════════════════════════════════════════════════════════════════════════
def foglio_pareggio(wb, d):
    ws = wb.create_sheet("Valore e pareggio")
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 34), ("B", 14), ("C", 14), ("D", 4), ("E", 58)):
        ws.column_dimensions[col].width = w
    v = d["valore"]

    scrivi(ws, "A1", "Quanto vale chi torna, e fino a quanto conviene spendere", T_TITOLO)
    scrivi(ws, "A2", "Spesa nei primi 365 giorni dal primo ordine, finestra uguale per "
                     "tutti.", T_TENUE)

    R0 = 4
    for i, t in enumerate(["", "clienti", "spesa media", "", "mediana"]):
        if t:
            scrivi(ws, f"{get_column_letter(1+i)}{R0}", t, T_TESTA, bordo=BORDO,
                   allinea=Alignment(horizontal="right" if i else "left"))
    scrivi(ws, f"A{R0+1}", "chi torna", T_FORTE)
    scrivi(ws, f"B{R0+1}", v["torna"]["clienti"], T_FORTE, "#,##0")
    scrivi(ws, f"C{R0+1}", v["torna"]["media"], T_FORTE, '#,##0 "£"')
    scrivi(ws, f"E{R0+1}", v["torna"]["mediana"], T_NORM, '#,##0 "£"')
    scrivi(ws, f"A{R0+2}", "chi si ferma al primo ordine", T_NORM)
    scrivi(ws, f"B{R0+2}", v["non_torna"]["clienti"], T_NORM, "#,##0")
    scrivi(ws, f"C{R0+2}", v["non_torna"]["media"], T_NORM, '#,##0 "£"')
    scrivi(ws, f"E{R0+2}", v["non_torna"]["mediana"], T_NORM, '#,##0 "£"')
    scrivi(ws, f"A{R0+4}", "Media e mediana sono lontane: la distribuzione ha una coda "
                           "lunga. Si mostrano tutte e due.", T_TENUE)

    r = R0 + 6
    scrivi(ws, f"A{r}", "MA CHI TORNA ERA GIÀ DIVERSO IL PRIMO GIORNO", T_H2); r += 1
    scrivi(ws, f"A{r}", "Se aveva speso di più al primo acquisto spenderà di più anche in "
                        "totale, e quella parte del divario non dipende dal ritorno. Si "
                        "confrontano allora solo clienti appaiati per dimensione del primo "
                        "ordine e mese di ingresso.", T_TENUE); r += 2

    for i, t in enumerate(["confronto", "differenza", "intervallo 95%"]):
        scrivi(ws, f"{get_column_letter(1+i)}{r}", t, T_TESTA, bordo=BORDO,
               allinea=Alignment(horizontal="right" if i else "left"))
    scrivi(ws, f"A{r+1}", "grezzo", T_NORM)
    scrivi(ws, f"B{r+1}", v["grezza"], T_NORM, '#,##0 "£"')
    scrivi(ws, f"C{r+1}", f"{v['grezza_ic'][0]:,} – {v['grezza_ic'][1]:,}".replace(",", "."),
           T_NORM, allinea=Alignment(horizontal="right"))
    scrivi(ws, f"A{r+2}", "fra clienti appaiati", T_FORTE)
    dif = scrivi(ws, f"B{r+2}", v["appaiata"], T_FORTE, '#,##0 "£"')
    scrivi(ws, f"C{r+2}", f"{v['appaiata_ic'][0]:,} – {v['appaiata_ic'][1]:,}".replace(",", "."),
           T_NORM, allinea=Alignment(horizontal="right"))
    r += 4
    scrivi(ws, f"A{r}", f"Sopravvive l'{round(v['appaiata']/v['grezza']*100)}% della "
                        f"differenza grezza. Non si dimezza: la dimensione del primo ordine "
                        f"spiega una fetta del divario, e non la fetta grossa.", T_TENUE)

    # ── il pareggio, con formule vere ────────────────────────────────────
    r += 3
    scrivi(ws, f"A{r}", "IL PUNTO DI PAREGGIO", T_H2); r += 1
    scrivi(ws, f"A{r}", "Cambia il margine nella cella gialla: la soglia si ricalcola.",
           T_NORM); r += 2

    r_marg = r
    scrivi(ws, f"A{r}", "Il tuo margine", T_FORTE)
    cel = scrivi(ws, f"B{r}", 0.20, Font(name="Calibri", size=12, bold=True, color=INCHIOSTRO),
                 "0%", riemp=INPUT, allinea=Alignment(horizontal="center"), bordo=BOX)
    scrivi(ws, f"E{r}", "← scrivi qui il tuo margine, fra 0% e 100%", T_TENUE)
    dv = DataValidation(type="decimal", operator="between", formula1=0, formula2=1,
                        allow_blank=False, showErrorMessage=True,
                        errorTitle="Margine fuori scala",
                        error="Il margine è una quota fra 0% e 100%.")
    ws.add_data_validation(dv)
    dv.add(cel)
    r += 2

    r_stima = r
    scrivi(ws, f"A{r}", "Differenza appaiata (stima)", T_NORM)
    scrivi(ws, f"B{r}", v["appaiata"], T_NORM, '#,##0 "£"'); r += 1
    r_basso = r
    scrivi(ws, f"A{r}", "Estremo basso dell'intervallo", T_NORM)
    scrivi(ws, f"B{r}", v["appaiata_ic"][0], T_NORM, '#,##0 "£"'); r += 2

    scrivi(ws, f"A{r}", "Soglia sulla stima", T_NORM)
    scrivi(ws, f"B{r}", f"=B{r_marg}*B{r_stima}", T_NORM, '#,##0 "£"')
    scrivi(ws, f"E{r}", "quanto rende in media riportare un cliente al secondo ordine",
           T_TENUE); r += 1

    scrivi(ws, f"A{r}", "Soglia prudente", T_FORTE)
    scrivi(ws, f"B{r}", f"=B{r_marg}*B{r_basso}", T_GROSSO, '#,##0 "£"')
    ws.row_dimensions[r].height = 28
    scrivi(ws, f"E{r}", "LA CIFRA SU CUI DECIDERE: regge anche se la stima è ottimista",
           T_FORTE); r += 3

    testo(ws, r, [
        "Perché una soglia e non un verdetto. La domanda di partenza era se convenga "
        "trattenere o acquisire: questi dati non possono dirlo, perché non contengono il "
        "costo di acquisizione né il costo del venduto.",
        "Ma il buco non si tappa dichiarandolo — si chiude cambiando cosa si consegna. "
        "Chi decide mette il proprio margine e legge la sua risposta, e il conto regge "
        "anche quando i suoi numeri cambiano.",
        "",
        "Il margine è un parametro dichiarato, non un risultato: senza il costo del "
        "venduto quello che si misura è ricavo, non profitto.",
    ], font=T_TENUE)


# ═════════════════════════════════════════════════════════════════════════
def foglio_dati(wb, d):
    """La matrice in formato lungo: una riga per cella. E' il foglio da cui si
    fa una tabella pivot senza dover disfare una matrice larga."""
    ws = wb.create_sheet("Dati coorti")
    intestazioni = ["coorte", "anno coorte", "mese coorte", "clienti entrati",
                    "mese relativo", "mese calendario", "anno calendario", "quota attivi"]
    for i, t in enumerate(intestazioni):
        scrivi(ws, f"{get_column_letter(1+i)}1", t, T_TESTA, bordo=BORDO)
        ws.column_dimensions[get_column_letter(1 + i)].width = max(12, len(t) + 3)
    r = 2
    for c in d["coorti"]:
        for i, val in enumerate(c["celle"]):
            if val is None:
                continue
            m = (c["mese"] - 1 + i) % 12 + 1
            a = c["anno"] + (c["mese"] - 1 + i) // 12
            for j, v in enumerate([c["coorte"], c["anno"], c["mese"], c["entrati"],
                                   i, MESI[m - 1], a]):
                scrivi(ws, f"{get_column_letter(1+j)}{r}", v, T_NORM)
            scrivi(ws, f"H{r}", val / 100, T_NORM, "0,0%")
            r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{r-1}"
    return r - 2


# ═════════════════════════════════════════════════════════════════════════
def verifica(d) -> list[str]:
    """Riapre il file appena scritto e lo confronta con i dati di partenza.

    Serve perche' openpyxl scrive senza lamentarsi anche quando il risultato e'
    sbagliato: una cella spostata di una riga, una formula diventata testo, una
    matrice allineata male. Qui si riapre e si guarda.
    """
    from openpyxl import load_workbook
    guai = []
    wb = load_workbook(USCITA)

    attesi = ["Leggimi", "Matrice per età", "Matrice per calendario",
              "Riacquisto", "Valore e pareggio", "Dati coorti"]
    for nome in attesi:
        if nome not in wb.sheetnames:
            guai.append(f"manca il foglio «{nome}»")
    if guai:
        return guai

    # le due matrici devono contenere le STESSE celle, solo in posizioni diverse
    c0 = d["coorti"][0]
    off = lambda c: (c["anno"] - c0["anno"]) * 12 + (c["mese"] - c0["mese"])
    for per_cal, nome in ((False, "Matrice per età"), (True, "Matrice per calendario")):
        ws = wb[nome]
        for j, c in enumerate(d["coorti"]):
            base = off(c) if per_cal else 0
            for i, val in enumerate(c["celle"]):
                if val is None:
                    continue
                letto = ws.cell(row=6 + j, column=3 + base + i).value
                if letto is None or abs(letto * 100 - val) > 0.051:
                    guai.append(f"{nome}: coorte {c['coorte']} mese +{i} "
                                f"vale {letto} invece di {val / 100}")
                    break

    # nel foglio per calendario, ogni colonna deve essere UN solo mese vero
    ws = wb["Matrice per calendario"]
    colonne = {}
    for j, c in enumerate(d["coorti"]):
        for i, val in enumerate(c["celle"]):
            if val is None:
                continue
            k = off(c) + i
            mese = ((c["mese"] - 1 + i) % 12 + 1, c["anno"] + (c["mese"] - 1 + i) // 12)
            colonne.setdefault(k, set()).add(mese)
    miste = [k for k, v in colonne.items() if len(v) > 1]
    if miste:
        guai.append(f"colonne che mescolano mesi diversi: {miste}")

    # il pareggio deve contenere FORMULE, non numeri incollati
    ws = wb["Valore e pareggio"]
    formule = [c.value for riga in ws.iter_rows() for c in riga
               if isinstance(c.value, str) and c.value.startswith("=")]
    if len(formule) < 2:
        guai.append(f"il foglio del pareggio ha {len(formule)} formule invece di 2: "
                    f"i valori sono incollati, e cambiare il margine non farebbe niente")

    # le cifre chiave devono comparire da qualche parte
    testi = " ".join(str(c.value) for ws in wb for riga in ws.iter_rows()
                     for c in riga if c.value is not None)
    v = d["valore"]
    for nome, atteso in (("differenza appaiata", v["appaiata"]),
                         ("estremo basso", v["appaiata_ic"][0]),
                         ("media di chi torna", v["torna"]["media"])):
        if str(atteso) not in testi and f"{atteso:,}".replace(",", ".") not in testi:
            guai.append(f"la cifra «{nome}» ({atteso}) non compare nel file")
    return guai


def main() -> None:
    if not DATI.is_file():
        raise SystemExit("Manca sito/dati.json. Lancia prima 08_dati_pagina.py.")
    d = json.loads(DATI.read_text(encoding="utf-8"))

    wb = Workbook()
    foglio_leggimi(wb, d)
    foglio_matrice(wb, d, per_calendario=False)
    foglio_matrice(wb, d, per_calendario=True)
    foglio_riacquisto(wb, d)
    foglio_pareggio(wb, d)
    n = foglio_dati(wb, d)

    wb.properties.title = "Clienti che tornano"
    wb.properties.subject = "Coorti, retention e riacquisto su Online Retail II"
    wb.properties.creator = "Raffaele Ciccone"
    wb.save(USCITA)

    print(f"scritto {USCITA.name}  ({USCITA.stat().st_size / 1024:.0f} KB)")
    print(f"  fogli: {', '.join(wb.sheetnames)}")
    print(f"  righe nel foglio dati: {n}")
    print(f"  differenza appaiata: {d['valore']['appaiata']:,} "
          f"[{d['valore']['appaiata_ic'][0]:,}, {d['valore']['appaiata_ic'][1]:,}]")

    guai = verifica(d)
    if guai:
        print(f"\nPROBLEMI ({len(guai)}):")
        for g in guai:
            print(f"  - {g}")
        sys.exit(1)
    print("  riaperto e ricontrollato: le due matrici contengono le stesse celle,")
    print("  le colonne del calendario sono mesi veri, il pareggio ha formule vive")


if __name__ == "__main__":
    main()
